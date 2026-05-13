#!/usr/bin/env python3
"""天猫BI仪表盘数据更新脚本 v6.0 - 修复广告数据汇总"""
import os, json, re, csv
from datetime import datetime
from collections import defaultdict
import openpyxl
from qcloud_cos import CosConfig, CosS3Client

print("="*50)
print("🦐 天猫BI数据解析 v6.0")
print(datetime.now().strftime("⏰ %Y-%m-%d %H:%M:%S"))
print("="*50)

# ── COS 连接 ──────────────────────────
secret_id  = os.environ['TENCENT_SECRET_ID']
secret_key = os.environ['TENCENT_SECRET_KEY']
bucket     = os.environ['TENCENT_COS_BUCKET']
region     = os.environ.get('TENCENT_COS_REGION', 'ap-beijing')

config = CosConfig(Region=region, SecretId=secret_id, SecretKey=secret_key)
client = CosS3Client(config)

# ── SKU 映射 ──────────────────────────
SKU_MAP = {
    'PET500_873480929689':      'PET500',
    'PET600_1001231224168':    'PET600',
    'RX400_Pro_704193543906':  'RX400_Pro',
    'U8_1032758801866':        'U8',
    'RX600_PRO_801617527631':  'RX600_PRO',
    'RX600P_800794914500':     'RX600P',
    'RX600_PROH_802250146018': 'RX600_PROH',
    '7232Pro_898077474925':   '7232Pro',
}

def norm_date(s):
    """兼容字符串、datetime对象、Excel数字日期"""
    from datetime import datetime as dt, timedelta
    if isinstance(s, dt):
        return s.strftime('%Y-%m-%d')
    s = str(s).strip()
    # Excel数字日期
    try:
        n = float(s)
        if 40000 < n < 50000:
            base = dt(1899, 12, 30)
            return (base + timedelta(days=n)).strftime('%Y-%m-%d')
    except ValueError:
        pass
    # 标准格式: 2026-04-28
    m = re.match(r'(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    # YYYYMMDD格式: 20260428
    m2 = re.match(r'(\d{4})(\d{2})(\d{2})', s)
    if m2:
        return f"{m2.group(1)}-{m2.group(2)}-{m2.group(3)}"
    return None

def pnum(s):
    """转数字，兜底0"""
    try:
        return float(str(s).replace(',', ''))
    except:
        return 0.0

def guess_ctype(fname, subdir):
    """根据文件名和子文件夹名判断数据类型"""
    fname_up = fname.upper()
    subdir_up = subdir.upper()
    
    if '退款' in fname_up or '退款' in subdir_up:
        return 'refund'
    if '流量' in fname_up or '流量' in subdir_up:
        return 'traffic'
    if '投放' in subdir_up or '推广' in fname_up:
        return 'ads'
    return 'sales'

# ── 下载 & 解析 ──────────────────────────
os.makedirs('cos-downloads', exist_ok=True)
all_data = {}   # {sku: {'sales':{dt:{}}, 'ads':{dt:{}}, 'refund':{dt:{}}, 'traffic':{dt:{}}}}
all_dates = set()

for sku_name, js_key in SKU_MAP.items():
    all_data[sku_name] = {'sales':{}, 'ads':{}, 'refund':{}, 'traffic':{}}
    prefix = f'data/{sku_name}/'
    
    try:
        # 列出所有子文件夹
        resp = client.list_objects(Bucket=bucket, Prefix=prefix, Delimiter='/', MaxKeys=100)
        subdirs = [p['Prefix'] for p in resp.get('CommonPrefixes', [])]
        
        if not subdirs:
            subdirs = [prefix]
        
        print(f"\n📦 [{sku_name}] {len(subdirs)}个子文件夹")
        
        for sub_prefix in subdirs:
            sub_name = sub_prefix.rstrip('/').split('/')[-1]
            
            # 列出该文件夹下所有文件
            resp2 = client.list_objects(Bucket=bucket, Prefix=sub_prefix, MaxKeys=100)
            files = [f['Key'] for f in resp2.get('Contents', []) if not f['Key'].endswith('/')]
            
            for key in files:
                fname = os.path.basename(key)
                if fname == '.gitkeep':
                    continue
                
                lpath = f'cos-downloads/{sku_name}/{fname}'
                os.makedirs(os.path.dirname(lpath), exist_ok=True)
                
                # 下载（流式读取）
                try:
                    r = client.get_object(Bucket=bucket, Key=key)
                    data = b''
                    chunk = r['Body'].read(8192)
                    while chunk:
                        data += chunk
                        chunk = r['Body'].read(8192)
                    
                    with open(lpath, 'wb') as f:
                        f.write(data)
                    
                except Exception as e:
                    print(f"  ❌ 下载失败 {fname}: {e}")
                    continue
                
                # 判断数据类型
                ct = guess_ctype(fname, sub_name)
                print(f"  📊 {sub_name}/{fname} → {ct}")
                
                # 解析Excel
                if fname.endswith(('.xlsx', '.xls')):
                    try:
                        wb = openpyxl.load_workbook(lpath, data_only=True)
                        for sheet_name in wb.sheetnames:
                            ws = wb[sheet_name]
                            rows = list(ws.iter_rows(values_only=True))
                            
                            # 找表头行
                            header_idx = None
                            for i, row in enumerate(rows):
                                if row and any('日期' in str(c) for c in row if c):
                                    header_idx = i
                                    break
                            if header_idx is None:
                                print(f"      ⚠️ 没找到表头")
                                continue
                            
                            headers = [str(c).strip() if c else '' for c in rows[header_idx]]
                            
                            for row in rows[header_idx+1:]:
                                if not row or not any(row):
                                    continue
                                
                                rd = dict(zip(headers, row))
                                
                                # 找日期列
                                dc = next((h for h in headers if '日期' in h or '时间' in h or 'date' in h.lower()), None)
                                if not dc:
                                    print(f"      ⚠️ 没找到日期列，表头={headers[:5]}")
                                    continue
                                dt = norm_date(rd.get(dc, ''))
                                if not dt:
                                    continue
                                
                                # 存入all_data
                                if ct == 'ads':
                                    if dt not in all_data[sku_name][ct]:
                                        all_data[sku_name][ct][dt] = {'花费': 0, '直接成交金额': 0, '间接成交金额': 0, '总成交金额': 0}
                                    # 累加每天所有计划的花费
                                    all_data[sku_name][ct][dt]['花费'] += pnum(rd.get('花费', 0))
                                    all_data[sku_name][ct][dt]['直接成交金额'] += pnum(rd.get('直接成交金额', 0))
                                    all_data[sku_name][ct][dt]['间接成交金额'] += pnum(rd.get('间接成交金额', 0))
                                    all_data[sku_name][ct][dt]['总成交金额'] += pnum(rd.get('总成交金额', 0))
                                elif ct == 'traffic':
                                    # 流量数据按日期累加访客数
                                    if dt not in all_data[sku_name][ct]:
                                        all_data[sku_name][ct][dt] = {'访客数': 0, '支付买家数': 0, '支付金额': 0}
                                    all_data[sku_name][ct][dt]['访客数'] += pnum(rd.get('访客数', 0))
                                    all_data[sku_name][ct][dt]['支付买家数'] += pnum(rd.get('支付买家数', 0))
                                    all_data[sku_name][ct][dt]['支付金额'] += pnum(rd.get('支付金额', 0))
                                else:
                                    all_data[sku_name][ct][dt] = rd
                                
                                all_dates.add(dt)
                        
                    except Exception as e:
                        print(f"  ❌ 解析错误 {fname}: {e}")
                        
                elif fname.endswith('.csv'):
                    try:
                        for enc in ['utf-8-sig', 'gbk', 'utf-8']:
                            try:
                                with open(lpath, 'r', encoding=enc) as f:
                                    reader = csv.DictReader(f)
                                    headers = reader.fieldnames
                                    dc = next((h for h in headers if '日期' in h), None)
                                    if not dc:
                                        continue
                                    for row in reader:
                                        dt = norm_date(row.get(dc, ''))
                                        if dt and dt != 'nan':
                                            if ct == 'refund':
                                                all_data[sku_name]['refund'][dt] = row
                                            else:
                                                all_data[sku_name]['sales'][dt] = row
                                            all_dates.add(dt)
                                    print(f'      ✅ CSV解析成功 ({enc})')
                                break
                            except:
                                continue
                    except Exception as e:
                        print(f"  ❌ CSV解析错误 {fname}: {e}")
                        
    except Exception as e:
        print(f"❌ [{sku_name}] 处理失败: {e}")

# ── 整理日期 ──────────────────────────────
dates_sorted = sorted(d for d in all_dates if d and d != 'nan' and len(d)==10 and '-' in d)
recent_14 = dates_sorted[-14:] if len(dates_sorted) > 14 else dates_sorted
print(f"\n📅 范围: {dates_sorted[0] if dates_sorted else '无'} ~ {dates_sorted[-1] if dates_sorted else '无'} | 使用 {len(recent_14)} 天")
print(f"DEBUG: all_dates长度={len(all_dates)}, 前5个={sorted(all_dates)[:5]}")

# ── 生成 data.json ───────────────────────
skus_output = {}

for sku_name, js_key in SKU_MAP.items():
    sku_info = all_data.get(sku_name, {})
    
    gmv, net, adSpend, adRev = [], [], [], []
    dir_, indir, roi, refund = [], [], [], []
    paidTraf, advTraf = [0.0]*len(recent_14), [0.0]*len(recent_14)
    revisit, inner = [0.0]*len(recent_14), [0.0]*len(recent_14)
    
    for idx, dt in enumerate(recent_14):
        sd = sku_info.get('sales', {}).get(dt, {})
        ad = sku_info.get('ads', {}).get(dt, {})
        rd = sku_info.get('refund', {}).get(dt, {})
        
        # 销售数据
        g = pnum(sd.get('支付金额', sd.get('销售额', 0))) / 10000
        r = pnum(sd.get('退款额', 0)) / 10000
        gmv.append(round(g, 2))
        net.append(round(g - r, 2))
        
        # 广告数据
        cost = pnum(ad.get('花费', 0)) / 10000
        dgmv = pnum(ad.get('直接成交金额', 0)) / 10000
        igmv = pnum(ad.get('间接成交金额', 0)) / 10000
        tgmv = pnum(ad.get('总成交金额', 0)) / 10000
        adSpend.append(round(cost, 2))
        adRev.append(round(tgmv, 2))
        dir_.append(round(dgmv, 2))
        indir.append(round(igmv, 2))
        roi.append(round(tgmv / cost, 2) if cost > 0 else 0)
        
        # 流量数据（从traffic字段读取）
        td = sku_info.get('traffic', {}).get(dt, {})
        visitors = pnum(td.get('访客数', 0))
        if visitors > 0:
            paidTraf[idx] = visitors
        rr_raw = str(rd.get('退款率', '0')).replace('%','')
        try:
            rv = float(rr_raw)
            rr = rv / 100 if rv > 1 else rv
        except:
            rr = 0
        refund.append(round(rr * 100, 2))
    
    # 检查数据是否有效
    has_gmv = any(x > 0 for x in gmv)
    has_ad = any(x > 0 for x in adSpend)
    
    if has_gmv:
        skus_output[js_key] = {
            'name': sku_name,
            'gmv': gmv,
            'net': net,
            'adSpend': adSpend,
            'adRev': adRev,
            'dir': dir_,
            'indir': indir,
            'roi': roi,
            'refund': refund,
            'paidTraf': paidTraf,
            'advTraf': advTraf,
            'revisit': revisit,
            'inner': inner,
        }
        print(f"✅ {js_key}: GMV={sum(gmv):.1f}万 | 广告花费={sum(adSpend):.1f}万 {'✅' if has_ad else '⚠️ 无广告数据'}")
    else:
        print(f"⚠️ {js_key}: 无有效销售数据")

# ── 写入 data.json ───────────────────────
output = {
    "dates": recent_14,
    "skus": skus_output,
    "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M")
}

with open('data.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print(f"\n🎉 更新完成! 生成 data.json ({len(recent_14)}天, {len(skus_output)}个SKU)")
print(f"   仪表盘: https://wg887k4tf4-lang.github.io/tmall-bi-dashboard/")
